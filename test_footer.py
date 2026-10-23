from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def add_page_number_to_run(run):
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    run._r.append(fldSimple)

doc = Document()
sec = doc.sections[0]
sec.left_margin = Inches(1.5)
sec.right_margin = Inches(1.0)
sec.top_margin = Inches(1.0)
sec.bottom_margin = Inches(1.0)

header = sec.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr1 = hp.add_run("Design | ")
hr1.font.name = 'Calibri'
hr1.font.size = Pt(11)
hr1.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

hr2 = hp.add_run("2026 - 2027")
hr2.font.name = 'Calibri'
hr2.font.size = Pt(11)
hr2.font.bold = True
hr2.font.color.rgb = RGBColor(0x1e, 0x3a, 0x8a)

pPr = hp._p.get_or_add_pPr()
pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:bottom w:val="single" w:sz="8" w:space="4" w:color="334155"/></w:pBdr>')
pPr.append(pBdr)

footer = sec.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
pPr_f = fp._p.get_or_add_pPr()
pBdr_f = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:top w:val="single" w:sz="8" w:space="6" w:color="94A3B8"/></w:pBdr>')
pPr_f.append(pBdr_f)

r_left = fp.add_run("AI Research Paper Assistant — Automated Academic Authoring Platform")
r_left.font.name = 'Calibri'
r_left.font.size = Pt(9.5)
r_left.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

# Tab to right
r_tab = fp.add_run("\t\t")

# Right badge: shaded maroon background run
r_badge = fp.add_run("  Page ")
r_badge.font.name = 'Calibri'
r_badge.font.size = Pt(10)
r_badge.font.bold = True
r_badge.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
rPr = r_badge._r.get_or_add_rPr()
shd = parse_xml(r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="clear" w:color="auto" w:fill="800000"/>')
rPr.append(shd)

r_num = fp.add_run()
r_num.font.name = 'Calibri'
r_num.font.size = Pt(10)
r_num.font.bold = True
r_num.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
rPr_num = r_num._r.get_or_add_rPr()
shd_num = parse_xml(r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="clear" w:color="auto" w:fill="800000"/>')
rPr_num.append(shd_num)
add_page_number_to_run(r_num)

r_end = fp.add_run("  ")
r_end.font.name = 'Calibri'
r_end.font.size = Pt(10)
r_end.font.bold = True
r_end.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
rPr_end = r_end._r.get_or_add_rPr()
shd_end = parse_xml(r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="clear" w:color="auto" w:fill="800000"/>')
rPr_end.append(shd_end)

doc.add_paragraph("Testing Header and Footer with Maroon Badge!")
doc.save("test_output.docx")
print("Saved test_output.docx successfully!")
