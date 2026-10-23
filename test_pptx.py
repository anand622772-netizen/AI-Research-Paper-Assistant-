import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

blank_slide_layout = prs.slide_layouts[6] # blank
slide = prs.slides.add_slide(blank_slide_layout)

# Add a test shape
shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(4), Inches(2))
shape.fill.solid()
shape.fill.fore_color.rgb = RGBColor(15, 23, 42)
shape.line.fill.background()

tf = shape.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Hello World"
p.font.size = Pt(24)
p.font.color.rgb = RGBColor(255, 255, 255)

prs.save("test_run.pptx")
print("Saved test_run.pptx successfully!")
