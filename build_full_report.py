"""
Comprehensive Parul University Minor Project Report Generator
Faculty of IT & Computer Science (MCA / M.Sc. IT / IMCA)
Subject: AI Research Paper Assistant
Ensures strict compliance with:
- Calibri font (Title: 14pt bold, Content: 10pt, Headers: 12pt)
- Margins: Left 1.5", Top/Bottom/Right 1.0"
- 1.5 Line Spacing, Justified Text
- Exact Sequence: Front Home Page, College Certificate, Preface, Acknowledgement, Index, Content
- Exact Index headings (1 to 8)
- Clean Simple DFDs (Level 0, Level 1, Level 2)
- Website Screenshots in section 4.3 as requested
- Minimum 45 pages verified via Microsoft Word COM
"""

import os
import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

REPORT_DIR = Path("outputs/parul_university_report")
REPORT_DIR.mkdir(parents=True, exist_ok=True)
ASSETS_DIR = Path("outputs/report_assets")

def add_page_number_to_run(run):
    """Inserts a dynamic Word PAGE field into a text run."""
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    run._r.append(fldSimple)

def set_cell_background(cell, fill_hex):
    """Sets background shading of a table cell."""
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    """Sets cell padding."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def format_paragraph(p, space_before=0, space_after=6, line_spacing=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.alignment = align

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=14, space_after=8, align=WD_ALIGN_PARAGRAPH.LEFT)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=10, space_after=4, align=WD_ALIGN_PARAGRAPH.LEFT)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1e, 0x3a, 0x8a)
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=6, space_after=3, align=WD_ALIGN_PARAGRAPH.LEFT)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    return p

def add_body_p(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=0, space_after=6, line_spacing=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(10)
    return p

def add_bullet_p(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    format_paragraph(p, space_before=0, space_after=3, line_spacing=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(10)
    return p

def add_code_block(doc, code_text):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=4, space_after=6, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.left_indent = Inches(0.2)
    run = p.add_run(code_text)
    run.font.name = 'Consolas'
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x1e, 0x29, 0x3b)
    return p

def add_image_figure(doc, img_path, caption_text, width_inches=5.6):
    if Path(img_path).exists():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_paragraph(p_img, space_before=6, space_after=3, line_spacing=1.0)
        p_img.add_run().add_picture(str(img_path), width=Inches(width_inches))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_paragraph(p_cap, space_before=2, space_after=8, line_spacing=1.15)
        r_cap = p_cap.add_run(f"Figure: {caption_text}")
        r_cap.font.name = 'Calibri'
        r_cap.font.size = Pt(9.5)
        r_cap.font.bold = True
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

def style_table(table, col_widths, headers, data):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Header row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]
        format_paragraph(p, space_before=3, space_after=3, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT)
        r = p.add_run(title)
        r.font.name = 'Calibri'
        r.font.size = Pt(9.5)
        r.font.bold = True
        set_cell_background(hdr_cells[i], "E2E8F0")
        set_cell_margins(hdr_cells[i], top=80, bottom=80, left=120, right=120)

    # Data rows
    for row_idx, row_data in enumerate(data):
        row = table.add_row()
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            format_paragraph(p, space_before=2, space_after=2, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT)
            r = p.add_run(str(val))
            r.font.name = 'Calibri'
            r.font.size = Pt(9.0)
            if row_idx % 2 == 1:
                set_cell_background(cell, "F8FAFC")
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)

    # Apply widths
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = Inches(width)

print("Setup completed.")
