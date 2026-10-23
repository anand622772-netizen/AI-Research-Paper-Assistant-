"""
Parul University Minor Project Report Generator
Faculty of IT & Computer Science (MCA / M.Sc. IT / IMCA)
Generates an exhaustive, 45+ page Microsoft Word (.docx) document
strictly complying with all institutional formatting guidelines and the provided index.
"""

import os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

REPORT_DIR = Path("outputs/parul_university_report")
REPORT_DIR.mkdir(parents=True, exist_ok=True)
ASSETS_DIR = Path("outputs/report_assets")

def add_page_number_to_run(run):
    """Inserts a dynamic Word PAGE field into a text run."""
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    run._r.append(fldSimple)

def set_cell_background(cell, fill_hex):
    """Sets background shading of a table cell."""
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    """Sets cell padding."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def format_paragraph(p, space_before=0, space_after=6, line_spacing=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.alignment = align

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=14, space_after=8, align=WD_ALIGN_PARAGRAPH.LEFT)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=10, space_after=4, align=WD_ALIGN_PARAGRAPH.LEFT)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1e, 0x3a, 0x8a)
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=6, space_after=3, align=WD_ALIGN_PARAGRAPH.LEFT)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    return p

def add_body_p(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=0, space_after=6, line_spacing=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(10)
    return p

def add_bullet_p(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    format_paragraph(p, space_before=0, space_after=3, line_spacing=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(10)
    return p

def add_code_block(doc, code_text):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=4, space_after=6, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.left_indent = Inches(0.2)
    run = p.add_run(code_text)
    run.font.name = 'Consolas'
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x1e, 0x29, 0x3b)
    return p

def add_image_figure(doc, img_path, caption_text, width_inches=5.8):
    if Path(img_path).exists():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_paragraph(p_img, space_before=6, space_after=3, line_spacing=1.0)
        p_img.add_run().add_picture(str(img_path), width=Inches(width_inches))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_paragraph(p_cap, space_before=2, space_after=8, line_spacing=1.15)
        r_cap = p_cap.add_run(f"Figure: {caption_text}")
        r_cap.font.name = 'Calibri'
        r_cap.font.size = Pt(9.5)
        r_cap.font.bold = True
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

def setup_section_margins_and_headers(section, chapter_title=""):
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1.0)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False

    # Header: Chapter title and year (Top Right), font 12
    hp = section.header.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    format_paragraph(hp, space_before=0, space_after=0, line_spacing=1.0)
    if chapter_title:
        hr = hp.add_run(f"{chapter_title} | 2026")
        hr.font.name = 'Calibri'
        hr.font.size = Pt(12)
        hr.font.color.rgb = RGBColor(0x64, 0x74, 0x8b)

    # Footer: Project title (Bottom Left) and page number (Bottom Right), font 10
    fp = section.footer.paragraphs[0]
    fp.text = ""
    format_paragraph(fp, space_before=0, space_after=0, line_spacing=1.0, align=WD_ALIGN_PARAGRAPH.LEFT)
    fr1 = fp.add_run("Minor Project: AI Research Paper Assistant")
    fr1.font.name = 'Calibri'
    fr1.font.size = Pt(10)
    fr1.font.color.rgb = RGBColor(0x64, 0x74, 0x8b)
    
    fr_tab = fp.add_run("\t\tPage ")
    fr_tab.font.name = 'Calibri'
    fr_tab.font.size = Pt(10)
    fr_tab.font.color.rgb = RGBColor(0x64, 0x74, 0x8b)
    add_page_number_to_run(fp.add_run())

print("Helper functions defined successfully.")
